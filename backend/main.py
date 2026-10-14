import os
import secrets
from datetime import date, datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, FastAPI, HTTPException, Query, Security
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from fastapi.security import APIKeyHeader

from sqlalchemy import case, func, select
from sqlalchemy.orm import Session

from database import Base, engine, get_db
from models import Transaction
from schemas import TransactionCreate, TransactionResponse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

IST = ZoneInfo("Asia/Kolkata")

# --------------------------------------------------
# API KEY AUTH
# --------------------------------------------------

API_KEY = os.environ["API_KEY"]  # set in Render's environment variables
api_key_header = APIKeyHeader(name="X-API-Key")


def require_key(key: str = Security(api_key_header)):
    if not secrets.compare_digest(key, API_KEY):
        raise HTTPException(status_code=401, detail="Invalid API key")


# --------------------------------------------------
# CREATE DATABASE TABLES
# --------------------------------------------------

Base.metadata.create_all(bind=engine)


# --------------------------------------------------
# FASTAPI
# --------------------------------------------------

app = FastAPI(
    title="Day To Day API",
    version="1.0.0",
)

# Every data route lives on this router, so it requires the API key.
# "/" and "/api/health" stay on `app` and remain public for Render's health check.
api = APIRouter(dependencies=[Depends(require_key)])


# --------------------------------------------------
# CORS
# --------------------------------------------------

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "https://daytwoday.netlify.app",
        # "http://localhost:3000",  # uncomment for local frontend testing
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
    expose_headers=["Content-Disposition"],
)


# --------------------------------------------------
# HEALTH CHECK (public)
# --------------------------------------------------

@app.get("/")
def home():
    return {"message": "Hello Welcome"}


@app.get("/api/health")
def health():
    return {"status": "ok"}


# --------------------------------------------------
# CREATE TRANSACTION
# --------------------------------------------------

@api.post(
    "/transactions",
    response_model=TransactionResponse,
)
def create_transaction(
    transaction: TransactionCreate,
    db: Session = Depends(get_db),
):
    new_transaction = Transaction(
        TranDate=transaction.TranDate,
        Description=transaction.Description,
        Amount=transaction.Amount,
        Type=transaction.Type,
        Mode=transaction.Mode,
        Remarks=transaction.Remarks,
        Other=transaction.Other,
        CategoryName=transaction.CategoryName,
    )

    db.add(new_transaction)
    db.commit()
    db.refresh(new_transaction)

    return new_transaction


# --------------------------------------------------
# GET ALL TRANSACTIONS
# --------------------------------------------------

@api.get(
    "/transactions",
    response_model=list[TransactionResponse],
)
def get_transactions(
    db: Session = Depends(get_db),
):
    statement = (
        select(Transaction)
        .where(Transaction.IsDeleted.is_(False))
        .order_by(
            Transaction.TranDate.desc(),
            Transaction.id.desc(),
        )
    )

    return db.scalars(statement).all()


# --------------------------------------------------
# GET TRANSACTIONS BY DATE
# --------------------------------------------------

@api.get(
    "/transactions/search",
    response_model=list[TransactionResponse],
)
def search_transactions(
    from_date: date = Query(...),
    to_date: date = Query(...),
    db: Session = Depends(get_db),
):
    if from_date > to_date:
        raise HTTPException(
            status_code=400,
            detail="from_date cannot be after to_date",
        )

    statement = (
        select(Transaction)
        .where(
            Transaction.IsDeleted.is_(False),
            Transaction.TranDate >= from_date,
            Transaction.TranDate <= to_date,
        )
        .order_by(
            Transaction.TranDate.desc(),
            Transaction.id.desc(),
        )
    )

    return db.scalars(statement).all()


# --------------------------------------------------
# SUMMARY (totals calculated in the database)
# --------------------------------------------------

@api.get("/transactions/summary")
def transaction_summary(
    from_date: date | None = Query(default=None),
    to_date: date | None = Query(default=None),
    db: Session = Depends(get_db),
):
    """Return active-transaction totals, optionally restricted to a date range."""
    if from_date and to_date and from_date > to_date:
        raise HTTPException(
            status_code=400,
            detail="from_date cannot be after to_date",
        )

    conditions = [Transaction.IsDeleted.is_(False)]
    if from_date:
        conditions.append(Transaction.TranDate >= from_date)
    if to_date:
        conditions.append(Transaction.TranDate <= to_date)

    type_norm = func.lower(func.trim(Transaction.Type))

    statement = select(
        func.count(Transaction.id),
        func.coalesce(
            func.sum(case((type_norm == "income", Transaction.Amount), else_=0)),
            0,
        ),
        func.coalesce(
            func.sum(case((type_norm == "expense", Transaction.Amount), else_=0)),
            0,
        ),
    ).where(*conditions)

    count, income, expenses = db.execute(statement).one()
    total_income = float(income)
    total_expenses = float(expenses)

    return {
        "from_date": from_date.isoformat() if from_date else None,
        "to_date": to_date.isoformat() if to_date else None,
        "transaction_count": count,
        "total_income": round(total_income, 2),
        "total_expenses": round(total_expenses, 2),
        "net_balance": round(total_income - total_expenses, 2),
    }


# --------------------------------------------------
# EXPORT TO EXCEL
# --------------------------------------------------

@api.get("/transactions/export")
def export_transactions(db: Session = Depends(get_db)):
    # Fetch ALL transactions (including soft-deleted), regardless of any UI filter.
    transactions = db.scalars(select(Transaction)).all()

    # Oldest first, so the export matches the Flutter app's backup order.
    chronological_transactions = sorted(
        transactions,
        key=lambda t: (t.TranDate, t.id),
    )

    # Running balance in chronological order.
    # Soft-deleted transactions do not affect the balance.
    running_balance = 0.0
    balance_by_id = {}

    for transaction in chronological_transactions:
        amount = float(transaction.Amount or 0)

        if not transaction.IsDeleted:
            transaction_type = (transaction.Type or "").strip().lower()

            if transaction_type == "income":
                running_balance += amount
            elif transaction_type == "expense":
                running_balance -= amount

        balance_by_id[transaction.id] = round(running_balance, 2)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Transactions"

    headers = [
        "Tran No",
        "TranDate",
        "Description",
        "Amount",
        "Type",
        "Mode",
        "Remarks",
        "Other",
        "CategoryName",
        "CreationDate",
        "IsDeleted",
        "Balance",
    ]
    worksheet.append(headers)

    # Format the header row.
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="1565C0")
        cell.alignment = Alignment(horizontal="center")

    # Add all database records.
    for transaction in chronological_transactions:
        # openpyxl cannot write timezone-aware datetimes.
        created = transaction.CreationDate
        if isinstance(created, datetime) and created.tzinfo is not None:
            created = created.astimezone(IST).replace(tzinfo=None)

        worksheet.append([
            transaction.id,
            transaction.TranDate,
            transaction.Description,
            transaction.Amount,
            transaction.Type,
            transaction.Mode,
            transaction.Remarks,
            transaction.Other,
            transaction.CategoryName,
            created,
            bool(transaction.IsDeleted),
            balance_by_id[transaction.id],
        ])

    # Column widths.
    widths = [12, 15, 32, 15, 15, 15, 25, 25, 20, 23, 15, 16]
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    # Generate the workbook in memory.
    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    # Backup filename with an IST timestamp.
    backup_filename = (
        f"Transactions_backup_till_"
        f"{datetime.now(IST).strftime('%Y-%m-%d_%H-%M-%S')}.xlsx"
    )

    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": f'attachment; filename="{backup_filename}"'
        },
    )


# --------------------------------------------------
# UPDATE TRANSACTION
# (keep {transaction_id} routes AFTER search/summary/export)
# --------------------------------------------------

@api.put(
    "/transactions/{transaction_id}",
    response_model=TransactionResponse,
)
def update_transaction(
    transaction_id: int,
    transaction: TransactionCreate,
    db: Session = Depends(get_db),
):
    existing_transaction = db.get(Transaction, transaction_id)

    if existing_transaction is None or existing_transaction.IsDeleted:
        raise HTTPException(
            status_code=404,
            detail="Transaction not found",
        )

    existing_transaction.TranDate = transaction.TranDate
    existing_transaction.Description = transaction.Description
    existing_transaction.Amount = transaction.Amount
    existing_transaction.Type = transaction.Type
    existing_transaction.Mode = transaction.Mode
    existing_transaction.Remarks = transaction.Remarks
    existing_transaction.Other = transaction.Other
    existing_transaction.CategoryName = transaction.CategoryName

    # Preserve id, CreationDate, and IsDeleted.
    db.commit()
    db.refresh(existing_transaction)

    return existing_transaction


# --------------------------------------------------
# SOFT DELETE
# --------------------------------------------------

@api.delete("/transactions/{transaction_id}")
def soft_delete_transaction(
    transaction_id: int,
    db: Session = Depends(get_db),
):
    transaction = db.get(Transaction, transaction_id)

    if transaction is None or transaction.IsDeleted:
        raise HTTPException(
            status_code=404,
            detail="Transaction not found",
        )

    transaction.IsDeleted = True
    db.commit()

    return {
        "message": "Transaction moved to deleted records",
        "id": transaction.id,
    }


# Register the protected routes (must come after all @api routes are defined).
app.include_router(api)
